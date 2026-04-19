"""综合导出 Excel（原 `绿联APP评论.py`）。

所有硬编码写入 openpyxl 的逻辑原样保留，仅把 `data.xxx` 的引用改为 `markets.xxx`。
"""
from __future__ import annotations

try:
    import _bootstrap  # noqa: F401
except ImportError:
    pass

from datetime import datetime
from typing import List, Optional

from openpyxl import Workbook

from common.logger import get_logger

from . import huawei, oppo, vivo, xiaomi

_logger = get_logger(__name__)


class Comment:
    def __init__(self) -> None:
        self.wb = Workbook()
        dt = datetime.now().strftime("%Y-%m-%d")
        self.filename = f"{dt}评论.xlsx"
        if "Sheet" in self.wb.sheetnames:
            self.wb.remove(self.wb["Sheet"])

    @staticmethod
    def _date(ts_millis: int) -> str:
        return datetime.fromtimestamp(ts_millis / 1000).strftime("%Y-%m-%d")

    # ---------- 小米 ----------

    def xiaomi_comment(self) -> None:
        data = xiaomi.getContent()
        rows = [["时间", "用户", "内容", "评论星", "机型", "版本", "版本号"]]
        if data and data.get("data", {}).get("commentInfoList"):
            for comment in data["data"]["commentInfoList"]:
                main = comment.get("main") or {}
                rows.append(
                    [
                        self._date(main["updateTime"]),
                        main.get("userName"),
                        main.get("content"),
                        main.get("score"),
                        main.get("device"),
                        main.get("versionName"),
                        main.get("versionCode"),
                    ]
                )
        self._save(rows, "小米评论", 1)

    def xiaomi_score(self) -> List:
        score = (xiaomi.getCore().get("data") or {}).get("score") or {}
        return [
            "小米",
            score.get("score"),
            score.get("count"),
            score.get("count5"),
            score.get("count4"),
            score.get("count3"),
            score.get("count2"),
            score.get("count1"),
        ]

    # ---------- vivo ----------

    def vivo_comment(self) -> None:
        payload = vivo.getContent()
        rows = [["时间", "用户", "评论星", "内容", "机型", "APP版本", "点赞数"]]
        for row in (payload.get("data") or {}).get("data") or []:
            rows.append(
                [
                    self._date(row["commentOn"]),
                    row.get("userName"),
                    row.get("score"),
                    row.get("comment"),
                    row.get("model"),
                    row.get("appVersion"),
                    row.get("goodCount"),
                ]
            )
        self._save(rows, "VIVO评论", 2)

    def vivo_score(self) -> Optional[List]:
        data = vivo.getCore().get("data") or {}
        if not data:
            return None
        return ["VIVO", round(data.get("starRating") or 0, 1), data.get("commentCount")]

    # ---------- OPPO ----------

    def oppo_comment(self) -> None:
        payload = oppo.getContent()
        rows = [["时间", "用户", "评分", "内容", "机型"]]
        for row in (payload.get("data") or {}).get("rows") or []:
            rows.append(
                [
                    row.get("UPDATE_TIME"),
                    row.get("USER_NICKNAME"),
                    row.get("USER_GRADE"),
                    row.get("PM_WORD"),
                    row.get("MOBILE_NAME"),
                ]
            )
        self._save(rows, "OPPO评论", 3)

    def oppo_score(self) -> List:
        data = oppo.getCore().get("data") or {}
        detail = data.get("score_detail") or {}
        return [
            "OPPO",
            round(data.get("agv_score") or 0, 1),
            data.get("count"),
            (detail.get("5") or {}).get("count"),
            (detail.get("4") or {}).get("count"),
            (detail.get("3") or {}).get("count"),
            (detail.get("2") or {}).get("count"),
            (detail.get("1") or {}).get("count"),
        ]

    # ---------- 华为 ----------

    def huawei_comment(self) -> None:
        payload = huawei.getContent()
        rows = [["时间", "用户", "评分", "内容", "机型", "版本"]]
        for row in ((payload.get("data") or {}).get("reviewList") or []):
            rows.append(
                [
                    self._date(row["operTimeStamp"]),
                    row.get("nickName"),
                    row.get("rating"),
                    row.get("content"),
                    row.get("phoneType"),
                    row.get("version"),
                ]
            )
        self._save(rows, "huawei评论", 4)

    def huawei_score(self) -> List:
        data = huawei.getCore().get("data") or {}
        return ["华为", round(data.get("star") or 0, 1), data.get("totalReviews"), None, None, None, None, None]

    # ---------- Excel 公共方法 ----------

    def _save(self, data: List[List], sheet_name: str, index: int) -> None:
        ws = self.wb.create_sheet(sheet_name, index)
        for row in data:
            ws.append(row)
        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = max_length + 2

    def save_score_sheet(self) -> None:
        rows: List[List] = [["平台", "评分", "评论数", "5星", "4星", "3星", "2星", "1星"]]
        rows.append(self.xiaomi_score())
        vivo_row = self.vivo_score()
        if vivo_row:
            rows.append(vivo_row)
        rows.append(self.oppo_score())
        rows.append(self.huawei_score())
        self._save(rows, "总评分", 0)

    def export(self) -> None:
        self.xiaomi_comment()
        self.vivo_comment()
        self.oppo_comment()
        self.huawei_comment()
        self.save_score_sheet()
        self.wb.save(self.filename)
        _logger.info("已导出 %s", self.filename)


if __name__ == "__main__":
    Comment().export()
